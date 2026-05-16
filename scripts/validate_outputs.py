from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from common_utils import ROOT

CLEANED_CSV = ROOT / 'cleaned_data' / 'cleaned_auction_properties.csv'
AI_CSV = ROOT / 'cleaned_data' / 'ai_property_reviews.csv'
FILTERED_XLSX = ROOT / 'output_excel' / 'filtered_properties_3000_to_8000.xlsx'
RANKED_XLSX = ROOT / 'output_excel' / 'investment_ranked_properties.xlsx'
MANUAL_XLSX = ROOT / 'output_excel' / 'manual_review_top_candidates.xlsx'

REQUIRED_CLEANED_COLUMNS = {
    'parcel_id', 'owner_name', 'property_address', 'city', 'state', 'zip_code', 'legal_description',
    'bid_cost', 'property_type', 'source_page', 'raw_text', 'extraction_confidence'
}
REQUIRED_AI_COLUMNS = {
    'parcel_id', 'ai_review_score', 'investment_category', 'manual_review_flag_ai', 'recommendation_ai',
    'reasoning_summary_ai', 'key_risks_ai', 'missing_information', 'next_due_diligence_step_ai', 'confidence_level'
}
REQUIRED_RANKED_COLUMNS = {
    'parcel_id', 'property_address', 'city', 'state', 'zip_code', 'bid_cost', 'legal_description', 'property_type',
    'source_page', 'extraction_confidence', 'rule_based_score', 'rule_based_category', 'bid_price_signal',
    'address_quality_signal', 'property_clarity_signal', 'crime_risk_estimate', 'surrounding_value_signal',
    'neighborhood_growth_signal', 'data_confidence', 'ai_review_score', 'investment_category', 'manual_review_flag',
    'recommendation', 'reasoning_summary', 'key_risks', 'missing_information', 'next_due_diligence_step', 'confidence_level'
}


def inspect_excel(path: Path):
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), values_only=True))
    return {'sheet_names': wb.sheetnames, 'max_row': ws.max_row, 'max_column': ws.max_column, 'top_rows': rows}


def passfail(name: str, ok: bool):
    print(f"{'PASS' if ok else 'FAIL'} {name}")
    return ok


def main():
    overall_ok = True

    cleaned_exists = CLEANED_CSV.exists()
    overall_ok &= passfail('cleaned CSV exists', cleaned_exists)
    cleaned_df = pd.read_csv(CLEANED_CSV) if cleaned_exists else pd.DataFrame()
    cleaned_ok = cleaned_exists and len(cleaned_df) > 0 and REQUIRED_CLEANED_COLUMNS.issubset(set(cleaned_df.columns))
    overall_ok &= passfail('cleaned CSV', cleaned_ok)
    print('cleaned CSV row count:', len(cleaned_df))
    print('cleaned CSV column count:', len(cleaned_df.columns))
    print('cleaned CSV top 5 rows:')
    print(cleaned_df.head().to_dict('records'))

    ai_exists = AI_CSV.exists()
    overall_ok &= passfail('AI review CSV exists', ai_exists)
    ai_df = pd.read_csv(AI_CSV) if ai_exists else pd.DataFrame()
    ai_ok = ai_exists and len(ai_df) > 0 and REQUIRED_AI_COLUMNS.issubset(set(ai_df.columns))
    overall_ok &= passfail('AI review CSV', ai_ok)
    print('AI review CSV row count:', len(ai_df))
    print('AI review CSV column count:', len(ai_df.columns))
    print('AI review CSV top 5 rows:')
    print(ai_df.head().to_dict('records'))

    for label, path in [('filtered Excel', FILTERED_XLSX), ('investment Excel', RANKED_XLSX), ('manual review Excel', MANUAL_XLSX)]:
        exists = path.exists()
        overall_ok &= passfail(f'{label} exists', exists)
        if exists:
            info = inspect_excel(path)
            ok = info['max_row'] > 1 and info['max_column'] > 1
            overall_ok &= passfail(label, ok)
            print(f'{label} sheet names:', info['sheet_names'])
            print(f'{label} row count:', info['max_row'])
            print(f'{label} column count:', info['max_column'])
            print(f'{label} top 5 rows:', info['top_rows'])

    ranked_df = pd.read_excel(RANKED_XLSX) if RANKED_XLSX.exists() else pd.DataFrame()
    ranked_ok = not ranked_df.empty and REQUIRED_RANKED_COLUMNS.issubset(set(ranked_df.columns))
    overall_ok &= passfail('investment ranked required columns', ranked_ok)

    manual_df = pd.read_excel(MANUAL_XLSX) if MANUAL_XLSX.exists() else pd.DataFrame()
    manual_ok = not manual_df.empty
    overall_ok &= passfail('manual review workbook non-empty', manual_ok)

    git_ls = subprocess.run(['git', 'ls-files', '.env'], cwd=ROOT, capture_output=True, text=True)
    env_tracked = bool(git_ls.stdout.strip())
    overall_ok &= passfail('.env not tracked by git', not env_tracked)

    if not overall_ok:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
