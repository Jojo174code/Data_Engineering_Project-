from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]

ILLEGAL_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
ADDRESS_RE = re.compile(r"\d+\s+[A-Z0-9].*", re.IGNORECASE)
ZIP_RE = re.compile(r"\b(74\d{3})\b")


def clean_text(value) -> str:
    if value is None:
        return ""
    text = str(value)
    text = ILLEGAL_CHAR_RE.sub("", text)
    return text.strip()


def normalize_city(value: str) -> str:
    text = clean_text(value)
    if not text:
        return "Unknown"
    replacements = {
        "CITY OF TULSA": "Tulsa",
        "TOWN OF JENKS": "Jenks",
        "TOWN OF BIXBY": "Bixby",
        "CITY OF OWASSO": "Owasso",
        "CITY OF SAND SPRINGS": "Sand Springs",
        "CITY OF BROKEN ARROW": "Broken Arrow",
        "CITY OF COLLINSVILLE": "Collinsville",
        "CITY OF SKIATOOK": "Skiatook",
    }
    return replacements.get(text.upper(), text.title())


def parse_money(value):
    text = clean_text(value).replace("$", "").replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def looks_residential_address(address: str) -> bool:
    address = clean_text(address)
    if not address or address.upper() == "ADDRESS UNKNOWN":
        return False
    return bool(ADDRESS_RE.match(address))


def infer_extraction_confidence(parcel_id: str, address: str, legal_description: str, bid_cost) -> str:
    score = 0
    if clean_text(parcel_id):
        score += 1
    if looks_residential_address(address):
        score += 1
    if clean_text(legal_description):
        score += 1
    if bid_cost is not None:
        score += 1
    if score >= 4:
        return "High"
    if score >= 2:
        return "Medium"
    return "Low"


def write_dataframe_to_excel(df, path: Path, sheet_name: str, currency_columns: Iterable[str] | None = None,
                             wrap_columns: Iterable[str] | None = None,
                             category_fill_column: str | None = None,
                             category_fills: dict[str, str] | None = None):
    import pandas as pd

    path.parent.mkdir(parents=True, exist_ok=True)
    if df.empty:
        raise ValueError(f"Refusing to write empty workbook: {path}")

    currency_columns = set(currency_columns or [])
    wrap_columns = set(wrap_columns or [])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.book[sheet_name]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        header_font = Font(bold=True)
        wrap_alignment = Alignment(wrap_text=True, vertical="top")

        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = wrap_alignment

        header_map = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                header = ws.cell(row=1, column=cell.column).value
                if header in currency_columns and isinstance(cell.value, (int, float)):
                    cell.number_format = '$#,##0.00'
                if header in wrap_columns:
                    cell.alignment = wrap_alignment

        if category_fill_column and category_fills and category_fill_column in header_map:
            col_idx = header_map[category_fill_column]
            for r in range(2, ws.max_row + 1):
                category = ws.cell(row=r, column=col_idx).value
                color = category_fills.get(str(category), None)
                if color:
                    fill = PatternFill(fill_type="solid", fgColor=color)
                    for c in range(1, ws.max_column + 1):
                        ws.cell(row=r, column=c).fill = fill

        for col_idx in range(1, ws.max_column + 1):
            max_len = 0
            col_letter = get_column_letter(col_idx)
            for cell in ws[col_letter]:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 50)

    validate_workbook(path, min_rows=2, min_cols=2)


def validate_workbook(path: Path, min_rows: int = 2, min_cols: int = 2):
    wb = load_workbook(path)
    sheet_names = wb.sheetnames
    if not sheet_names:
        raise ValueError(f"Workbook has no sheets: {path}")
    ws = wb[sheet_names[0]]
    if ws.max_row < min_rows:
        raise ValueError(f"Workbook has too few rows: {path} ({ws.max_row})")
    if ws.max_column < min_cols:
        raise ValueError(f"Workbook has too few columns: {path} ({ws.max_column})")
    return {
        "sheet_names": sheet_names,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "first_rows": [[cell.value for cell in row] for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row))],
    }
