from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
INPUT_FILE = BASE_DIR / "input" / "rc_assets_final_bid_list_original.xlsx"
CLEANED_CSV = BASE_DIR / "cleaned_data" / "rc_assets_final_cleaned_list.csv"
REMOVED_CSV = BASE_DIR / "cleaned_data" / "removed_red_strikethrough_properties.csv"
LOG_FILE = BASE_DIR / "logs" / "clean_log.txt"

SORT_PRIORITY_COLUMNS = [
    "Auction Order",
    "Auction List Order",
    "Original Auction Order",
    "Source Row",
    "Source Page",
    "Page",
    "Original #",
    "#",
    "Rank",
    "Original Rank",
]


@dataclass
class RowAudit:
    excel_row: int
    remove: bool
    color_value: str
    has_red_indicator: bool
    has_strikethrough: bool
    crossed_out_indicator: bool
    parcel_id: str
    reason: str


def normalize_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def looks_red(value: str) -> bool:
    return normalize_text(value).lower() == "red"


def cell_has_red_fill(cell) -> bool:
    fill = cell.fill
    if not fill or fill.patternType != "solid":
        return False
    rgb = getattr(fill.fgColor, "rgb", None)
    if not rgb:
        return False
    rgb = str(rgb).upper()
    redish = {"00FFC7CE", "00F4CCCC", "00FF0000", "00E06666", "00EA9999"}
    return rgb in redish or rgb.endswith("C7CE") or rgb.endswith("0000")


def row_has_strikethrough(row_cells: Iterable) -> bool:
    for cell in row_cells:
        if getattr(cell.font, "strike", False):
            return True
    return False


def row_has_crossed_out_indicator(values: list[str]) -> bool:
    joined = " | ".join(normalize_text(v).lower() for v in values if v is not None)
    indicators = [
        "crossed out",
        "crossed-out",
        "strikethrough",
        "strike through",
        "remove",
        "removed",
        "do not bid",
        "deleted",
    ]
    return any(token in joined for token in indicators)


def choose_sort_column(columns: list[str]) -> str | None:
    for candidate in SORT_PRIORITY_COLUMNS:
        if candidate in columns:
            return candidate
    return None


def main() -> int:
    BASE_DIR.joinpath("cleaned_data").mkdir(parents=True, exist_ok=True)
    BASE_DIR.joinpath("logs").mkdir(parents=True, exist_ok=True)

    wb = load_workbook(INPUT_FILE)
    ws = wb["Bid List"] if "Bid List" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows())
    headers = [normalize_text(cell.value) for cell in rows[0]]

    audits: list[RowAudit] = []
    data_rows: list[list] = []
    removed_rows: list[list] = []

    for row in rows[1:]:
        values = [cell.value for cell in row]
        row_map = dict(zip(headers, values))
        color_value = normalize_text(row_map.get("Color"))
        has_red_indicator = looks_red(color_value) or any(cell_has_red_fill(cell) for cell in row)
        has_strike = row_has_strikethrough(row)
        crossed_out_indicator = row_has_crossed_out_indicator(values)
        remove = has_red_indicator and (has_strike or crossed_out_indicator)
        parcel_id = normalize_text(row_map.get("Parcel ID"))
        reason = "red+crossed-out" if remove else "kept"
        audit = RowAudit(
            excel_row=row[0].row,
            remove=remove,
            color_value=color_value,
            has_red_indicator=has_red_indicator,
            has_strikethrough=has_strike,
            crossed_out_indicator=crossed_out_indicator,
            parcel_id=parcel_id,
            reason=reason,
        )
        audits.append(audit)
        augmented = values + [
            row[0].row,
            audit.has_red_indicator,
            audit.has_strikethrough,
            audit.crossed_out_indicator,
            audit.reason,
        ]
        if remove:
            removed_rows.append(augmented)
        else:
            data_rows.append(augmented)

    augmented_headers = headers + [
        "source_excel_row",
        "red_indicator_detected",
        "strikethrough_detected",
        "crossed_out_indicator_detected",
        "cleaning_action",
    ]

    active_df = pd.DataFrame(data_rows, columns=augmented_headers)
    removed_df = pd.DataFrame(removed_rows, columns=augmented_headers)

    sort_column = choose_sort_column(active_df.columns.tolist())
    if sort_column:
        numeric_sort = pd.to_numeric(active_df[sort_column], errors="coerce")
        active_df = active_df.assign(_sort=numeric_sort.fillna(active_df["source_excel_row"]))
        active_df = active_df.sort_values(["_sort", "source_excel_row"], kind="stable").drop(columns=["_sort"])
    else:
        active_df = active_df.sort_values(["source_excel_row"], kind="stable")

    if active_df.empty:
        raise SystemExit("Cleaned list has 0 rows, stopping.")

    active_df.to_csv(CLEANED_CSV, index=False, quoting=csv.QUOTE_MINIMAL)
    removed_df.to_csv(REMOVED_CSV, index=False, quoting=csv.QUOTE_MINIMAL)

    removed_parcels = [a.parcel_id for a in audits if a.remove and a.parcel_id]
    kept_count = len([a for a in audits if not a.remove and a.parcel_id])
    log_lines = [
        f"input_file={INPUT_FILE}",
        f"sheet_used={ws.title}",
        f"original_row_count={len(rows) - 1}",
        f"removed_red_strikethrough_row_count={len(removed_rows)}",
        f"final_active_row_count={len(active_df)}",
        f"removed_parcel_ids={removed_parcels}",
        f"kept_parcel_ids_count={kept_count}",
        f"columns_preserved={headers}",
        f"sort_column_used={sort_column or 'source_excel_row'}",
    ]
    LOG_FILE.write_text("\n".join(log_lines) + "\n", encoding="utf-8")
    print("\n".join(log_lines))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
