from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parents[1]
CLEANED_CSV = BASE_DIR / "cleaned_data" / "rc_assets_final_cleaned_list.csv"
REMOVED_CSV = BASE_DIR / "cleaned_data" / "removed_red_strikethrough_properties.csv"
CRIME_CSV = BASE_DIR / "cleaned_data" / "crime_scan_results.csv"
ECON_CSV = BASE_DIR / "cleaned_data" / "economic_scan_results.csv"
AI_CSV = BASE_DIR / "cleaned_data" / "ai_scanner_results.csv"
OUTPUT_FILE = BASE_DIR / "output_excel" / "rc_assets_final_scanned_bid_list.xlsx"
CHEAPEST_FILE = BASE_DIR / "output_excel" / "rc_assets_final_scanned_by_auction_order.xlsx"
REMOVED_XLSX = BASE_DIR / "output_excel" / "rc_assets_removed_red_strikethrough.xlsx"

AI_FILL = {
    "Top Candidate": "00C6EFCE",
    "Strong Candidate": "00E2F0D9",
    "Watchlist": "00FFF2CC",
    "High Risk": "00F4CCCC",
    "Remove": "00FFC7CE",
}
COLOR_FILL = {
    "green": "00E2F0D9",
    "blue": "00DDEBF7",
    "purple": "00E4DFEC",
    "yellow": "00FFF2CC",
    "red": "00F4CCCC",
}
MAX_BID_FILL = PatternFill(fill_type="solid", fgColor="00FFF2CC")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="00D9EAD3")
THIN_BORDER = Border(
    left=Side(style="thin", color="00D9D9D9"),
    right=Side(style="thin", color="00D9D9D9"),
    top=Side(style="thin", color="00D9D9D9"),
    bottom=Side(style="thin", color="00D9D9D9"),
)

SORT_COLUMNS = [
    "Auction Order",
    "Auction List Order",
    "Original Auction Order",
    "Original Rank",
    "Source Row",
    "Source Page",
    "Page",
    "Original #",
    "#",
    "Rank",
    "source_excel_row",
]


def safe_read_csv(path: Path, fallback_columns: list[str] | None = None) -> pd.DataFrame:
    try:
        return pd.read_csv(path)
    except pd.errors.EmptyDataError:
        return pd.DataFrame(columns=fallback_columns or [])


def load_merged() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    cleaned = pd.read_csv(CLEANED_CSV)
    removed = safe_read_csv(REMOVED_CSV, cleaned.columns.tolist())
    crime = pd.read_csv(CRIME_CSV)
    econ = pd.read_csv(ECON_CSV)
    ai = pd.read_csv(AI_CSV)
    merged = cleaned.merge(crime, on=["Parcel ID", "Address"], how="left").merge(econ, on=["Parcel ID", "Address"], how="left").merge(ai, on=["Parcel ID", "Address"], how="left")
    return merged, removed, crime, econ, ai


def sort_main(df: pd.DataFrame) -> pd.DataFrame:
    for col in SORT_COLUMNS:
        if col in df.columns:
            return df.assign(_sort=pd.to_numeric(df[col], errors="coerce")).sort_values(["_sort", "source_excel_row"], kind="stable").drop(columns=["_sort"])
    return df.sort_values(["source_excel_row"], kind="stable")


def sort_cheapest(df: pd.DataFrame) -> pd.DataFrame:
    return df.assign(_bid=pd.to_numeric(df["Opening Bid"], errors="coerce")).sort_values(["_bid", "source_excel_row"], kind="stable").drop(columns=["_bid"])


def apply_base_style(cell):
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.border = THIN_BORDER


def apply_color_fill(cell, color_value: str):
    fill_rgb = COLOR_FILL.get((color_value or "").strip().lower())
    if fill_rgb:
        cell.fill = PatternFill(fill_type="solid", fgColor=fill_rgb)


def write_sheet(ws, df: pd.DataFrame):
    ws.append(df.columns.tolist())
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        apply_base_style(cell)

    color_idx = df.columns.get_loc("Color") + 1 if "Color" in df.columns else None
    max_bid_idx = df.columns.get_loc("Max Bid") + 1 if "Max Bid" in df.columns else None
    category_idx = df.columns.get_loc("ai_final_category") + 1 if "ai_final_category" in df.columns else None

    for _, row in df.iterrows():
        ws.append(row.tolist())
        current_row = ws.max_row
        color_value = str(row.get("Color", ""))
        for idx in range(1, ws.max_column + 1):
            cell = ws.cell(current_row, idx)
            apply_base_style(cell)
            if color_idx:
                apply_color_fill(cell, color_value)
        if category_idx:
            val = ws.cell(current_row, category_idx).value
            if val in AI_FILL:
                ws.cell(current_row, category_idx).fill = PatternFill(fill_type="solid", fgColor=AI_FILL[val])
        if max_bid_idx:
            ws.cell(current_row, max_bid_idx).fill = MAX_BID_FILL

    for currency_col in ["Opening Bid", "Max Bid", "Bid Spread"]:
        if currency_col in df.columns:
            idx = df.columns.get_loc(currency_col) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(r, idx).number_format = '$#,##0.00'

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx in range(1, ws.max_column + 1):
        width = max(len(str(ws.cell(r, col_idx).value or "")) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 12), 42)


def summary_frame(merged: pd.DataFrame, removed: pd.DataFrame) -> pd.DataFrame:
    rows = []
    rows.append(("original property count", len(merged) + len(removed)))
    rows.append(("removed red/strikethrough count", len(removed)))
    rows.append(("final active property count", len(merged)))
    for name, series in [
        ("Color", merged["Color"].value_counts(dropna=False)),
        ("Grade", merged["Grade"].value_counts(dropna=False) if "Grade" in merged.columns else pd.Series(dtype=int)),
        ("ai_final_category", merged["ai_final_category"].value_counts(dropna=False)),
        ("ai_bid_priority", merged["ai_bid_priority"].value_counts(dropna=False)),
        ("crime_risk_level", merged["crime_risk_level"].value_counts(dropna=False)),
        ("economic_signal_level", merged["economic_signal_level"].value_counts(dropna=False)),
        ("development_match_strength", merged["development_match_strength"].value_counts(dropna=False)),
    ]:
        for key, value in series.items():
            rows.append((f"{name}: {key}", int(value)))
    opening = pd.to_numeric(merged["Opening Bid"], errors="coerce")
    maxbid = pd.to_numeric(merged["Max Bid"], errors="coerce")
    rows.extend([
        ("total opening bid", float(opening.fillna(0).sum())),
        ("total max bid", float(maxbid.fillna(0).sum())),
        ("cheapest active property", merged.loc[opening.idxmin(), "Address"] if not merged.empty else ""),
        ("most expensive active property", merged.loc[opening.idxmax(), "Address"] if not merged.empty else ""),
        ("number of properties with blank max bid", int(maxbid.isna().sum())),
    ])
    return pd.DataFrame(rows, columns=["metric", "value"])


def main() -> int:
    merged, removed, crime, econ, ai = load_merged()
    merged = sort_main(merged)
    cheapest = sort_cheapest(merged.copy())
    top = merged[(merged["ai_final_category"].isin(["Top Candidate", "Strong Candidate"])) | (merged["ai_bid_priority"].isin(["Priority 1", "Priority 2"]))].copy()
    summary = summary_frame(merged, removed)

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for title, df in [
        ("Final List - Auction Order", merged),
        ("Final List - Cheapest First", cheapest),
        ("Top Candidates", top),
        ("Removed Red Crossed Out", removed),
        ("Crime Scan", crime),
        ("Economic Scan", econ),
        ("Summary", summary),
    ]:
        ws = wb.create_sheet(title)
        write_sheet(ws, df.fillna(""))

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    wb.save(CHEAPEST_FILE)

    removed_wb = Workbook()
    rws = removed_wb.active
    rws.title = "Removed Red Crossed Out"
    write_sheet(rws, removed.fillna(""))
    removed_wb.save(REMOVED_XLSX)
    print(f"saved={OUTPUT_FILE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
