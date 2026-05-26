from __future__ import annotations

import re
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

REPO_DIR = Path(__file__).resolve().parents[3]
BASE_DIR = Path(__file__).resolve().parents[1]
INPUT_WORKBOOK = BASE_DIR / "output_excel" / "rc_assets_final_scanned_bid_list.xlsx"
AUCTION_PDF = REPO_DIR / "2026 auction.pdf"
OUTPUT_WORKBOOK = BASE_DIR / "output_excel" / "rc_assets_final_scanned_bid_list_CORRECTED.xlsx"
OUTPUT_ACTIVE_CSV = BASE_DIR / "cleaned_data" / "rc_assets_final_cleaned_list_CORRECTED.csv"
OUTPUT_DELETED_CSV = BASE_DIR / "cleaned_data" / "deleted_properties.csv"
LOG_FILE = BASE_DIR / "logs" / "correction_log.txt"

DELETE_TARGETS = [
    ("44675-92-08-04460", "5134 W CHARLES PAGE"),
    ("24925-02-36-17420", "1084 N NORFOLK AV E"),
    ("05400-02-25-04330", "524 E YOUNG ST N"),
    ("44675-92-08-04450", "5128 W CHARLES PAGE"),
    ("34300-03-32-15230", "135 N COLUMBIA AV E"),
    ("34750-02-26-08160", "2507 N MAIN ST E TULSA"),
    ("07600-02-24-01200", "820 E 36 ST N"),
    ("46465-92-11-31880", "2034 S PHOENIX AV W"),
    ("00575-93-06-01170", "511 S VICTOR AV E"),
]
DELETE_IDS = {parcel for parcel, _ in DELETE_TARGETS}
MANUAL_AUCTION_OVERRIDES = {
    "53775-02-02-00220": (967, 143, 967),
    "53825-02-02-01690": (971, 144, 971),
    "90306-03-06-24020": (1249, 184, 1249),
    "96203-62-03-50220": (1303, 193, 1303),
}
CRIME_COLUMNS = [
    "Parcel ID",
    "Address",
    "crime_risk_level",
    "crime_context_summary",
    "crime_data_granularity",
    "crime_source_title",
    "crime_source_url",
    "crime_data_confidence",
]
ECON_COLUMNS = [
    "Parcel ID",
    "Address",
    "economic_signal_level",
    "development_match_strength",
    "matched_development_area",
    "matched_project_or_investment",
    "development_source_title",
    "development_source_url",
    "development_summary",
    "development_match_reason",
    "economic_data_confidence",
]
NEW_ORDER_COLUMNS = ["Auction List No", "Auction Source Page", "Auction Source Sequence"]
HEADER_FILL = PatternFill(fill_type="solid", fgColor="00D9EAD3")
ROW_COLOR_FILLS = {
    "green": PatternFill(fill_type="solid", fgColor="00E2F0D9"),
    "blue": PatternFill(fill_type="solid", fgColor="00DDEBF7"),
    "purple": PatternFill(fill_type="solid", fgColor="00E4DFEC"),
    "yellow": PatternFill(fill_type="solid", fgColor="00FFF2CC"),
    "red": PatternFill(fill_type="solid", fgColor="00F4CCCC"),
}
AI_CATEGORY_FILLS = {
    "Top Candidate": PatternFill(fill_type="solid", fgColor="00C6EFCE"),
    "Strong Candidate": PatternFill(fill_type="solid", fgColor="00D9EAD3"),
    "Watchlist": PatternFill(fill_type="solid", fgColor="00FFF2CC"),
    "High Risk": PatternFill(fill_type="solid", fgColor="00F4CCCC"),
    "Remove": PatternFill(fill_type="solid", fgColor="00EA9999"),
}
MAX_BID_FILL = PatternFill(fill_type="solid", fgColor="00FFE699")
HEADER_FILLS = {
    "original": PatternFill(fill_type="solid", fgColor="00D9EAD3"),
    "auction": PatternFill(fill_type="solid", fgColor="00DDEBF7"),
    "crime": PatternFill(fill_type="solid", fgColor="00F4CCCC"),
    "economic": PatternFill(fill_type="solid", fgColor="00E4DFEC"),
    "ai": PatternFill(fill_type="solid", fgColor="00FCE5CD"),
    "helper": PatternFill(fill_type="solid", fgColor="00EEEEEE"),
}
THIN_BORDER = Border(
    left=Side(style="thin", color="00D9D9D9"),
    right=Side(style="thin", color="00D9D9D9"),
    top=Side(style="thin", color="00D9D9D9"),
    bottom=Side(style="thin", color="00D9D9D9"),
)


def normalize_parcel(value) -> str:
    return str(value).strip()


def extract_auction_map(pdf_path: Path) -> dict[str, tuple[int, int, int]]:
    reader = PdfReader(str(pdf_path))
    pattern = re.compile(r"\b(\d{1,4})\s+\n\s*(\d{5}-\d{2}-\d{2}-\d{5})\b", re.M)
    order_map: dict[str, tuple[int, int, int]] = {}
    sequence = 0
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for match in pattern.finditer(text):
            prop_no = int(match.group(1))
            parcel = match.group(2)
            if parcel not in order_map:
                sequence += 1
                order_map[parcel] = (prop_no, page_no, sequence)
    order_map.update(MANUAL_AUCTION_OVERRIDES)
    return order_map


def load_base_frames() -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
    sheets = pd.read_excel(INPUT_WORKBOOK, sheet_name=None)
    main = sheets["Final List - Auction Order"].copy()
    return main, sheets


def build_style_map(workbook_path: Path):
    wb = load_workbook(workbook_path)
    ws = wb["Final List - Auction Order"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    style_map = {}
    for r in range(2, ws.max_row + 1):
        parcel = normalize_parcel(ws.cell(r, headers.index("Parcel ID") + 1).value)
        style_map[parcel] = {
            "fills": [copy(ws.cell(r, c).fill) for c in range(1, ws.max_column + 1)],
            "fonts": [copy(ws.cell(r, c).font) for c in range(1, ws.max_column + 1)],
            "borders": [copy(ws.cell(r, c).border) for c in range(1, ws.max_column + 1)],
            "alignments": [copy(ws.cell(r, c).alignment) for c in range(1, ws.max_column + 1)],
            "formats": [ws.cell(r, c).number_format for c in range(1, ws.max_column + 1)],
        }
    return style_map, headers


def add_auction_order_columns(df: pd.DataFrame, auction_map: dict[str, tuple[int, int, int]]) -> pd.DataFrame:
    df = df.copy()
    prop_nos = []
    pages = []
    sequences = []
    found_flags = []
    fallbacks = []
    for _, row in df.iterrows():
        parcel = normalize_parcel(row["Parcel ID"])
        match = auction_map.get(parcel)
        if match:
            prop_no, page_no, sequence = match
            found_flags.append(0)
            prop_nos.append(prop_no)
            pages.append(page_no)
            sequences.append(sequence)
        else:
            found_flags.append(1)
            prop_nos.append(pd.NA)
            pages.append(pd.NA)
            sequences.append(pd.NA)
        fallback = row.get("Original Rank")
        if pd.isna(fallback):
            fallback = row.get("source_excel_row")
        fallbacks.append(fallback)
    df["Auction List No"] = prop_nos
    df["Auction Source Page"] = pages
    df["Auction Source Sequence"] = sequences
    df["_auction_found"] = found_flags
    df["_fallback_order"] = pd.to_numeric(pd.Series(fallbacks), errors="coerce")
    return df


def sort_active(df: pd.DataFrame) -> pd.DataFrame:
    return df.sort_values(
        by=["_auction_found", "Auction List No", "Auction Source Page", "Auction Source Sequence", "_fallback_order", "source_excel_row"],
        kind="stable",
        na_position="last",
    ).drop(columns=["_auction_found", "_fallback_order"])


def sort_cheapest(df: pd.DataFrame) -> pd.DataFrame:
    return df.assign(_bid=pd.to_numeric(df["Opening Bid"], errors="coerce")).sort_values(["_bid", "source_excel_row"], kind="stable").drop(columns=["_bid"])


def classify_header(header: str) -> str:
    if header in {"Auction List No", "Auction Source Page", "Auction Source Sequence"}:
        return "auction"
    if header.startswith("crime_") or header in CRIME_COLUMNS:
        return "crime"
    if header.startswith("economic_") or header.startswith("development_") or header in ECON_COLUMNS or header in {"matched_development_area", "matched_project_or_investment"}:
        return "economic"
    if header.startswith("ai_"):
        return "ai"
    if header in {"source_excel_row", "red_indicator_detected", "strikethrough_detected", "crossed_out_indicator_detected", "cleaning_action", "deletion_reason"}:
        return "helper"
    return "original"


def apply_base_style(cell):
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.border = THIN_BORDER


def preferred_width(header: str, measured: int) -> int:
    wide = {
        "Address": 28,
        "Notes": 42,
        "crime_context_summary": 42,
        "development_summary": 42,
        "development_match_reason": 38,
        "ai_summary": 44,
        "ai_key_risks": 32,
        "ai_due_diligence_next_step": 34,
        "crime_source_url": 34,
        "development_source_url": 34,
        "Parcel ID": 19,
        "Auction List No": 15,
    }
    if header in wide:
        return max(wide[header], min(measured, 48))
    if header in {"Color", "Grade", "#", "Original Rank", "Auction Source Page", "Auction Source Sequence"}:
        return max(12, min(measured, 16))
    return max(14, min(measured, 26))


def write_df_sheet(ws, df: pd.DataFrame, style_map: dict, base_headers: list[str], hidden_columns: set[str] | None = None):
    hidden_columns = hidden_columns or set()
    ws.append(df.columns.tolist())
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(1, col)
        cell.font = Font(bold=True, size=12)
        cell.fill = HEADER_FILLS[classify_header(str(cell.value))]
        apply_base_style(cell)
    ws.row_dimensions[1].height = 28
    color_idx = df.columns.get_loc("Color") + 1 if "Color" in df.columns else None
    max_bid_idx = df.columns.get_loc("Max Bid") + 1 if "Max Bid" in df.columns else None
    ai_category_idx = df.columns.get_loc("ai_final_category") + 1 if "ai_final_category" in df.columns else None
    for _, row in df.iterrows():
        ws.append(row.tolist())
        current_row = ws.max_row
        parcel = normalize_parcel(row.get("Parcel ID")) if "Parcel ID" in df.columns else None
        styles = style_map.get(parcel)
        row_fill = ROW_COLOR_FILLS.get(str(row.get("Color", "")).strip().lower())
        for idx, header in enumerate(df.columns, start=1):
            cell = ws.cell(current_row, idx)
            apply_base_style(cell)
            cell.font = Font(size=11)
            if row_fill:
                cell.fill = copy(row_fill)
            if styles and header in base_headers:
                base_idx = base_headers.index(header)
                cell.font = copy(styles["fonts"][base_idx])
                cell.border = copy(styles["borders"][base_idx])
                cell.alignment = copy(styles["alignments"][base_idx])
                cell.number_format = styles["formats"][base_idx]
            if header in {"Notes", "crime_context_summary", "development_summary", "development_match_reason", "ai_summary", "ai_key_risks", "ai_due_diligence_next_step"}:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if max_bid_idx:
            ws.cell(current_row, max_bid_idx).fill = copy(MAX_BID_FILL)
        if ai_category_idx:
            cat = str(ws.cell(current_row, ai_category_idx).value)
            if cat in AI_CATEGORY_FILLS:
                ws.cell(current_row, ai_category_idx).fill = copy(AI_CATEGORY_FILLS[cat])
        for currency_col in ["Opening Bid", "Max Bid", "Bid Spread"]:
            if currency_col in df.columns:
                cidx = df.columns.get_loc(currency_col) + 1
                ws.cell(current_row, cidx).number_format = "$#,##0.00"
        long_text = max(len(str(row.get(col, "") or "")) for col in [c for c in ["Notes", "crime_context_summary", "development_summary", "ai_summary"] if c in df.columns]) if any(c in df.columns for c in ["Notes", "crime_context_summary", "development_summary", "ai_summary"]) else 0
        ws.row_dimensions[current_row].height = 42 if long_text > 120 else 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 110
    for col_idx in range(1, ws.max_column + 1):
        header = str(ws.cell(1, col_idx).value)
        width = max(len(str(ws.cell(r, col_idx).value or "")) for r in range(1, ws.max_row + 1)) + 2
        ws.column_dimensions[get_column_letter(col_idx)].width = preferred_width(header, width)
        if header in hidden_columns:
            ws.column_dimensions[get_column_letter(col_idx)].hidden = True


def order_columns(df: pd.DataFrame) -> pd.DataFrame:
    preferred = [
        "#", "Color", "Source", "Original Rank", "Auction List No", "Auction Source Page", "Auction Source Sequence",
        "Parcel ID", "Address", "Opening Bid", "Max Bid", "Bid Spread", "Grade", "Drive-By Status",
        "Lien/Title Risk", "Entitlement/Zoning Risk", "Final Decision", "Notes",
        "crime_risk_level", "crime_context_summary", "crime_data_granularity", "crime_source_title", "crime_source_url", "crime_data_confidence",
        "economic_signal_level", "development_match_strength", "matched_development_area", "matched_project_or_investment", "development_source_title", "development_source_url", "development_summary", "development_match_reason", "economic_data_confidence",
        "ai_final_score", "ai_final_category", "ai_bid_priority", "ai_summary", "ai_key_risks", "ai_due_diligence_next_step", "ai_confidence",
        "source_excel_row", "red_indicator_detected", "strikethrough_detected", "crossed_out_indicator_detected", "cleaning_action", "deletion_reason",
    ]
    ordered = [col for col in preferred if col in df.columns]
    ordered += [col for col in df.columns if col not in ordered]
    return df[ordered]


def make_summary(previous_count: int, active_df: pd.DataFrame, deleted_df: pd.DataFrame) -> pd.DataFrame:
    rows = [
        ("previous active property count", previous_count),
        ("deleted property count", len(deleted_df)),
        ("final active property count", len(active_df)),
    ]
    for name, series in [
        ("Color", active_df["Color"].value_counts(dropna=False)),
        ("Grade", active_df["Grade"].value_counts(dropna=False)),
        ("ai_final_category", active_df["ai_final_category"].value_counts(dropna=False)),
        ("ai_bid_priority", active_df["ai_bid_priority"].value_counts(dropna=False)),
    ]:
        for key, value in series.items():
            rows.append((f"{name}: {key}", int(value)))
    opening = pd.to_numeric(active_df["Opening Bid"], errors="coerce")
    maxbid = pd.to_numeric(active_df["Max Bid"], errors="coerce")
    rows.extend([
        ("total opening bid", float(opening.fillna(0).sum())),
        ("total max bid", float(maxbid.fillna(0).sum())),
        ("number of properties with blank max bid", int(maxbid.isna().sum())),
    ])
    return pd.DataFrame(rows, columns=["metric", "value"])


def main() -> int:
    auction_map = extract_auction_map(AUCTION_PDF)
    base_main, _ = load_base_frames()
    style_map, base_headers = build_style_map(INPUT_WORKBOOK)
    previous_count = len(base_main)
    base_main["Parcel ID"] = base_main["Parcel ID"].map(normalize_parcel)
    base_main["Address"] = base_main["Address"].astype(str)

    deleted_df = base_main[base_main["Parcel ID"].isin(DELETE_IDS)].copy()
    active_df = base_main[~base_main["Parcel ID"].isin(DELETE_IDS)].copy()
    if len(deleted_df) != len(DELETE_IDS):
        missing = sorted(DELETE_IDS - set(deleted_df["Parcel ID"]))
        raise SystemExit(f"Missing deletion targets in base workbook: {missing}")

    deleted_df["deletion_reason"] = "User-directed urgent deletion"
    active_df = add_auction_order_columns(active_df, auction_map)
    deleted_df = add_auction_order_columns(deleted_df, auction_map)
    active_df = sort_active(active_df)
    deleted_df = sort_active(deleted_df)
    active_df = order_columns(active_df)
    deleted_df = order_columns(deleted_df)
    cheapest_df = order_columns(sort_cheapest(active_df.copy()))
    top_df = active_df[
        active_df["ai_final_category"].isin(["Top Candidate", "Strong Candidate"]) |
        active_df["ai_bid_priority"].isin(["Priority 1", "Priority 2"])
    ].copy()
    crime_df = active_df[CRIME_COLUMNS].copy()
    econ_df = active_df[ECON_COLUMNS].copy()
    summary_df = make_summary(previous_count, active_df, deleted_df)

    OUTPUT_ACTIVE_CSV.parent.mkdir(parents=True, exist_ok=True)
    active_df.to_csv(OUTPUT_ACTIVE_CSV, index=False)
    deleted_df.to_csv(OUTPUT_DELETED_CSV, index=False)

    wb = Workbook()
    wb.remove(wb.active)
    sheets = [
        ("Final List - Auction Order", active_df),
        ("Cheapest First", cheapest_df),
        ("Top Candidates", top_df),
        ("Deleted Properties", deleted_df),
        ("Crime Scan", crime_df),
        ("Economic Scan", econ_df),
        ("Summary", summary_df),
    ]
    hidden_cols = {"source_excel_row", "red_indicator_detected", "strikethrough_detected", "crossed_out_indicator_detected", "cleaning_action"}
    for title, df in sheets:
        ws = wb.create_sheet(title)
        write_df_sheet(ws, df.fillna(""), style_map, base_headers, hidden_columns=hidden_cols if title in {"Final List - Auction Order", "Cheapest First", "Top Candidates", "Deleted Properties"} else set())
    OUTPUT_WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_WORKBOOK)

    log_lines = [
        f"input_workbook={INPUT_WORKBOOK}",
        f"auction_pdf={AUCTION_PDF}",
        f"previous_active_count={previous_count}",
        f"deleted_count={len(deleted_df)}",
        f"final_active_count={len(active_df)}",
        f"deleted_parcels={sorted(deleted_df['Parcel ID'].tolist())}",
        f"pdf_order_matches={active_df['Auction List No'].notna().sum()}",
        f"pdf_order_missing={active_df['Auction List No'].isna().sum()}",
        f"output_workbook={OUTPUT_WORKBOOK}",
    ]
    LOG_FILE.write_text("\n".join(log_lines) + "\n", encoding="utf-8")
    print("\n".join(log_lines))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
