#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from common_custom_33 import ALLOWED_CONFIDENCE, ALLOWED_RECOMMENDATIONS, ALLOWED_TIERS, CLEANED_DIR, CUSTOM_PROPERTIES, LOG_DIR, OUTPUT_DIR, clean_text, write_log

BASE_CSV = CLEANED_DIR / 'custom_33_base.csv'
AI_CSV = CLEANED_DIR / 'custom_33_ai_reviews.csv'
FINAL_CSV = CLEANED_DIR / 'custom_33_final_ranked.csv'
ENRICH_CSV = CLEANED_DIR / 'custom_33_enriched_from_prior_iterations.csv'
GEOCODE_CSV = CLEANED_DIR / 'custom_33_geocoded.csv'
AREA_CSV = CLEANED_DIR / 'custom_33_area_intelligence.csv'
LOG_FILE = LOG_DIR / 'validation_report.txt'
EXCEL_FILES = [
    OUTPUT_DIR / 'custom_33_property_investment_ranking.xlsx',
    OUTPUT_DIR / 'custom_top_10_priority_targets.xlsx',
    OUTPUT_DIR / 'custom_drive_by_list.xlsx',
    OUTPUT_DIR / 'custom_high_risk_watchlist.xlsx',
]


def check_reasoning(text: str, row: pd.Series) -> bool:
    lower = clean_text(text).lower()
    has_bid = any(token in lower for token in [str(row['bid_cost']).lower(), f"{float(row['bid_cost']):.2f}".lower()])
    has_property = clean_text(row['parcel_id']).lower() in lower or clean_text(row['property_address']).lower() in lower
    has_crime = clean_text(row['crime_risk_level']).lower() in lower
    has_econ = clean_text(row['economic_risk_level']).lower() in lower or any(token in lower for token in ['income', 'poverty', 'unemployment', 'economic'])
    return has_bid and has_property and has_crime and has_econ


def main() -> int:
    errors = []
    report = []
    expected_parcels = {item['parcel_id'] for item in CUSTOM_PROPERTIES}

    if not BASE_CSV.exists():
        errors.append('custom_33_base.csv missing')
        base = pd.DataFrame()
    else:
        base = pd.read_csv(BASE_CSV)
        if len(base) != 33:
            errors.append(f'custom_33_base.csv row count {len(base)} != 33')

    if not FINAL_CSV.exists():
        errors.append('custom_33_final_ranked.csv missing')
        final = pd.DataFrame()
    else:
        final = pd.read_csv(FINAL_CSV)
        if len(final) != 33:
            errors.append(f'custom_33_final_ranked.csv row count {len(final)} != 33')

    if not AI_CSV.exists():
        errors.append('custom_33_ai_reviews.csv missing')
        ai = pd.DataFrame()
    else:
        ai = pd.read_csv(AI_CSV)
        if len(ai) != 33:
            errors.append(f'custom_33_ai_reviews.csv row count {len(ai)} != 33')

    if not final.empty:
        dupes = final['parcel_id'][final['parcel_id'].duplicated()].tolist()
        if dupes:
            errors.append(f'duplicate parcel IDs in final output: {dupes}')
        missing_parcels = sorted(expected_parcels - set(final['parcel_id'].astype(str)))
        if missing_parcels:
            errors.append(f'missing parcel IDs in final output: {missing_parcels}')
        bad_tiers = sorted(set(final['final_tier']) - set(ALLOWED_TIERS))
        if bad_tiers:
            errors.append(f'invalid final_tier values: {bad_tiers}')
        bad_recs = sorted(set(final['recommendation']) - set(ALLOWED_RECOMMENDATIONS))
        if bad_recs:
            errors.append(f'invalid recommendation values: {bad_recs}')
        bad_conf = sorted(set(final['confidence_level']) - set(ALLOWED_CONFIDENCE))
        if bad_conf:
            errors.append(f'invalid confidence_level values: {bad_conf}')
        for _, row in final.iterrows():
            if not check_reasoning(row['reasoning_summary'], row):
                errors.append(f"reasoning summary missing required context for parcel {row['parcel_id']}")
                break

    geocoding_success_count = 0
    crime_coverage_count = 0
    economic_coverage_count = 0
    matched_prior = 0
    if ENRICH_CSV.exists():
        enrich = pd.read_csv(ENRICH_CSV)
        matched_prior = int((enrich['enrichment_match_method'] != 'unmatched').sum())
    if GEOCODE_CSV.exists():
        geocode = pd.read_csv(GEOCODE_CSV)
        geocoding_success_count = int((geocode['geocode_status'] != 'Failed').sum())
    if AREA_CSV.exists():
        area = pd.read_csv(AREA_CSV)
        crime_coverage_count = int((area['crime_source'] != 'Unknown').sum())
        economic_coverage_count = int((area['economic_source'] != 'Unknown').sum())

    for path in EXCEL_FILES:
        if not path.exists():
            errors.append(f'missing Excel file: {path.name}')
            continue
        try:
            wb = load_workbook(path)
            ws = wb.active
            if ws.max_row <= 1 or ws.max_column <= 1:
                errors.append(f'Excel file has no data rows: {path.name}')
        except Exception as exc:
            errors.append(f'failed to open {path.name}: {exc}')

    repo_root = Path(__file__).resolve().parents[3]
    gitignore_ok = True
    import subprocess
    status = subprocess.run(['git', 'status', '--short'], cwd=repo_root, capture_output=True, text=True)
    status_lines = status.stdout.splitlines()
    for line in status_lines:
        if '.env' in line:
            errors.append('.env is staged or modified in git status')
        if '.venv' in line:
            errors.append('.venv is staged or modified in git status')

    tier_counts = final['final_tier'].value_counts().to_dict() if not final.empty else {}
    top10_count = min(10, len(final)) if not final.empty else 0
    report.extend([
        f'total custom properties: {len(expected_parcels)}',
        f'matched from prior iterations: {matched_prior}',
        f'geocoding success count: {geocoding_success_count}',
        f'crime data coverage count: {crime_coverage_count}',
        f'economic data coverage count: {economic_coverage_count}',
        f"Tier 1 count: {tier_counts.get('Tier 1 Priority', 0)}",
        f"Tier 2 count: {tier_counts.get('Tier 2 Strong Review', 0)}",
        f"Tier 3 count: {tier_counts.get('Tier 3 Watch', 0)}",
        f"Tier 4 count: {tier_counts.get('Tier 4 High Risk', 0)}",
        f'top 10 count: {top10_count}',
    ])
    if errors:
        report.append('validation FAIL')
        report.extend(errors)
        write_log(LOG_FILE, report)
        print('validation=FAIL')
        return 1
    report.append('validation PASS')
    write_log(LOG_FILE, report)
    print('validation=PASS')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
