#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[3]
OUTPUT_XLSX = REPO_ROOT / 'output_excel' / 'rc_assets_final_bid_list.xlsx'
OUTPUT_CSV = REPO_ROOT / 'cleaned_data' / 'rc_assets_final_bid_list.csv'

ROWS = [
    ['Green','Manual Review',1,'00575-93-06-01170','511 S VICTOR AV E',3910.27],
    ['Green','Manual Review',2,'34300-03-32-15230','135 N COLUMBIA AV E',3005.69],
    ['Green','Manual Review',8,'44200-02-11-04390','4696 N BOULDER AV W',3075.02],
    ['Green','Manual Review',9,'34750-02-26-08160','2507 N MAIN ST E TULSA',3097.02],
    ['Green','Manual Review',11,'23175-03-30-13610','2526 N QUAKER AV E',3115.02],
    ['Green','Manual Review',15,'53775-02-02-00220','6105 N MAIN ST E TULSA',3140.32],
    ['Purple','Manual Review',16,'23125-02-25-15350','1864 N OWASSO AV E',3140.35],
    ['Green','Manual Review',17,'07600-02-24-01200','820 E 36 ST N',3152.35],
    ['Green','Manual Review',21,'44200-02-14-06630','18 E 44 PL N',3183.01],
    ['Green','Manual Review',26,'07675-02-13-01140','4018 N GARRISON AV E',3245.67],
    ['Green','Manual Review',34,'17675-02-24-07570','750 E 32 PL N',3310.33],
    ['Green','Manual Review',35,'25500-03-29-06060','2625 E TECUMSEH ST N',3310.33],
    ['Green','Manual Review',36,'44200-02-11-03380','4627 N BOSTON AV E',3311.00],
    ['Green','Manual Review',38,'25375-03-29-05240','3215 E XYLER ST N',3313.80],
    ['Purple','Manual Review',39,'90306-03-06-24020','1641 E 56 ST N',3314.33],
    ['Green','Manual Review',40,'44250-02-12-15190','549 E 54 PL N',3327.00],
    ['Green','Manual Review',43,'46465-92-11-31880','2034 S PHOENIX AV W',3329.00],
    ['Purple','Manual Review',48,'29625-02-36-21010','419 E NEWTON PL N',3388.99],
    ['Purple','Manual Review',49,'10450-03-19-50550','3128 N TROOST PL E',3389.11],
    ['Green','$1,500–$2,000',2,'47175-03-29-18430','1903 N FLORENCE PL E',1549.17],
    ['Green','$1,500–$2,000',7,'38675-93-10-06110','1321 S DARLINGTON AV E',1603.63],
    ['Green','$1,500–$2,000',12,'40250-03-31-14910','1945 E NEWTON ST N',1613.17],
    ['Purple','$1,500–$2,000',25,'53825-02-02-01690','6369 N BOULDER AV W',1693.83],
    ['Green','$1,500–$2,000',28,'40825-02-13-08310','4337 N GARRISON AV E',1697.16],
    ['Green','$1,500–$2,000',31,'24675-03-31-13270','1404 N ST LOUIS AV E',1697.83],
    ['Purple','$1,500–$2,000',36,'56925-92-20-02770','3930 S 61 AV W',1716.68],
    ['Blue','$1,500–$2,000',41,'13850-93-05-10570','2621 E ADMIRAL CT N',1759.82],
    ['Green','$1,500–$2,000',43,'40850-02-12-07900','517 E 47 PL N',1771.82],
    ['Green','$1,500–$2,000',44,'40250-03-31-15190','2143 E NEWTON ST N',1778.61],
    ['Green','$1,500–$2,000',46,'44675-92-08-04450','5128 W CHARLES PAGE',1783.15],
    ['Green','$1,000–$1,500',5,'24925-02-36-17420','1084 N NORFOLK AV E',1199.79],
    ['Green','$1,000–$1,500',13,'40875-02-13-11530','210 E 44 ST N',1184.93],
    ['Green','$1,000–$1,500',16,'26725-02-25-16640','515 E UTE ST N',1216.55],
    ['Green','$1,000–$1,500',21,'05400-02-25-04330','524 E YOUNG ST N',1418.58],
    ['Purple','$1,000–$1,500',28,'90212-02-12-26700','1118 E 50 PL N',1317.38],
    ['Green','$1,000–$1,500',32,'41025-02-01-04710','5901 N GARRISON PL E',1273.20],
    ['Green','$1,000–$1,500',33,'41025-02-01-05430','522 E 59 ST N',1273.20],
    ['Green','$1,000–$1,500',36,'43300-03-29-09940','2311 N ATLANTA AV E',1179.88],
    ['Green','$1,000–$1,500',38,'05850-03-30-04070','1844 N TRENTON AV E',1297.53],
    ['Blue','$1,000–$1,500',50,'96203-62-03-50220','18911 S 29 AV W',1064.76],
    ['Green','$900–$1,000',1,'01875-03-19-00200','2612 N TRENTON AV E',962.57],
    ['Green','$900–$1,000',7,'44675-92-08-04460','5134 W CHARLES PAGE',907.78],
]

COLUMNS = ['Color','Source','Original Rank','Parcel ID','Address','Opening Bid']
FINAL_COLUMNS = ['#','Color','Source','Original Rank','Parcel ID','Address','Opening Bid','Max Bid','Bid Spread','Drive-By Status','Lien/Title Risk','Entitlement/Zoning Risk','Final Decision','Notes']
COLOR_FILLS = {
    'Green': 'E2F0D9',
    'Purple': 'E4DFEC',
    'Blue': 'DDEBF7',
}
MAX_BID_FILL = 'FFF2CC'
DROPDOWNS = {
    'J': ['Not Checked','Looks Good','Needs Caution','Avoid'],
    'K': ['Not Checked','Low','Medium','High','Avoid'],
    'L': ['Not Checked','Low','Medium','High','Avoid'],
    'M': ['Research','Drive By','Bid Candidate','Watch','Remove'],
}


def build_df() -> pd.DataFrame:
    df = pd.DataFrame(ROWS, columns=COLUMNS)
    if len(df) != 42:
        raise ValueError(f'Expected 42 rows, found {len(df)}')
    if df['Parcel ID'].duplicated().any():
        dupes = df.loc[df['Parcel ID'].duplicated(), 'Parcel ID'].tolist()
        raise ValueError(f'Duplicate Parcel IDs found: {dupes}')
    df = df.sort_values(['Opening Bid', 'Original Rank', 'Parcel ID'], ascending=[True, True, True]).reset_index(drop=True)
    df.insert(0, '#', range(1, len(df) + 1))
    df['Max Bid'] = None
    df['Bid Spread'] = None
    df['Drive-By Status'] = 'Not Checked'
    df['Lien/Title Risk'] = 'Not Checked'
    df['Entitlement/Zoning Risk'] = 'Not Checked'
    df['Final Decision'] = 'Research'
    df['Notes'] = ''
    return df[FINAL_COLUMNS]


def add_data_validations(ws, last_row: int):
    for col, values in DROPDOWNS.items():
        formula = '"' + ','.join(values) + '"'
        dv = DataValidation(type='list', formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{last_row}')


def format_bid_list(ws, df: pd.DataFrame):
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    last_row = len(df) + 1
    add_data_validations(ws, last_row)
    for row_idx in range(2, last_row + 1):
        color = ws.cell(row_idx, headers['Color']).value
        fill = COLOR_FILLS.get(color)
        if fill:
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row_idx, col_idx).fill = PatternFill(fill_type='solid', start_color=fill, end_color=fill)
        ws.cell(row_idx, headers['Max Bid']).fill = PatternFill(fill_type='solid', start_color=MAX_BID_FILL, end_color=MAX_BID_FILL)
        ws.cell(row_idx, headers['Bid Spread']).value = f'=IF(H{row_idx}="","",H{row_idx}-G{row_idx})'
    for col_name in ['Opening Bid', 'Max Bid', 'Bid Spread']:
        col_idx = headers[col_name]
        for row_idx in range(2, last_row + 1):
            ws.cell(row_idx, col_idx).number_format = '$#,##0.00'
    notes_col = headers['Notes']
    for row_idx in range(2, last_row + 1):
        ws.cell(row_idx, notes_col).alignment = Alignment(wrap_text=True, vertical='top')
    for col_idx in range(1, ws.max_column + 1):
        letter = ws.cell(1, col_idx).column_letter
        max_len = max(len(str(ws.cell(r, col_idx).value or '')) for r in range(1, min(ws.max_row, 250) + 1))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 28)


def build_summary(wb, df: pd.DataFrame):
    ws = wb.create_sheet('Summary')
    ws['A1'] = 'Metric'
    ws['B1'] = 'Value'
    ws['A1'].font = ws['B1'].font = Font(bold=True)
    cheapest = df.iloc[0]
    most_expensive = df.iloc[-1]
    rows = [
        ('total properties', len(df)),
        ('total opening bid amount', '=SUM(\'Bid List\'!G2:G43)'),
        ('count by Color - Green', '=COUNTIF(\'Bid List\'!B2:B43,"Green")'),
        ('count by Color - Purple', '=COUNTIF(\'Bid List\'!B2:B43,"Purple")'),
        ('count by Color - Blue', '=COUNTIF(\'Bid List\'!B2:B43,"Blue")'),
        ('count by Source - Manual Review', '=COUNTIF(\'Bid List\'!C2:C43,"Manual Review")'),
        ('count by Source - $1,500–$2,000', '=COUNTIF(\'Bid List\'!C2:C43,"$1,500–$2,000")'),
        ('count by Source - $1,000–$1,500', '=COUNTIF(\'Bid List\'!C2:C43,"$1,000–$1,500")'),
        ('count by Source - $900–$1,000', '=COUNTIF(\'Bid List\'!C2:C43,"$900–$1,000")'),
        ('cheapest property', f"{cheapest['Parcel ID']} | {cheapest['Address']} | ${cheapest['Opening Bid']:.2f}"),
        ('most expensive property', f"{most_expensive['Parcel ID']} | {most_expensive['Address']} | ${most_expensive['Opening Bid']:.2f}"),
        ('blank Max Bid count', '=COUNTBLANK(\'Bid List\'!H2:H43)'),
        ('total Max Bid amount', '=SUM(\'Bid List\'!H2:H43)'),
        ('total Bid Spread', '=SUM(\'Bid List\'!I2:I43)'),
    ]
    for i, (metric, value) in enumerate(rows, start=2):
        ws.cell(i, 1, metric)
        ws.cell(i, 2, value)
    for row in range(2, len(rows) + 2):
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical='top')
    for cell_ref in ['B3', 'B13', 'B14']:
        ws[cell_ref].number_format = '$#,##0.00'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 42


def validate_workbook(path: Path, expected_df: pd.DataFrame) -> dict:
    wb = load_workbook(path)
    if 'Bid List' not in wb.sheetnames:
        raise ValueError('Bid List sheet missing')
    if 'Summary' not in wb.sheetnames:
        raise ValueError('Summary sheet missing')
    ws = wb['Bid List']
    property_rows = ws.max_row - 1
    if property_rows != 42:
        raise ValueError(f'Expected 42 property rows, found {property_rows}')
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    if 'Max Bid' not in headers:
        raise ValueError('Max Bid column missing')
    spread_col = headers.index('Bid Spread') + 1
    first_formula = ws.cell(2, spread_col).value
    if not isinstance(first_formula, str) or not first_formula.startswith('=IF('):
        raise ValueError('Bid Spread formula missing')
    opening_col = headers.index('Opening Bid') + 1
    opening_values = [ws.cell(r, opening_col).value for r in range(2, ws.max_row + 1)]
    if opening_values != sorted(opening_values):
        raise ValueError('Opening Bid sort validation failed')
    parcel_col = headers.index('Parcel ID') + 1
    parcels = [str(ws.cell(r, parcel_col).value) for r in range(2, ws.max_row + 1)]
    if len(parcels) != len(set(parcels)):
        raise ValueError('Duplicate Parcel IDs found in workbook')
    if set(parcels) != set(expected_df['Parcel ID'].astype(str)):
        raise ValueError('Workbook parcel coverage mismatch')
    return {
        'property_rows': property_rows,
        'sorted': True,
        'duplicates_ok': True,
    }


def main():
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    df = build_df()
    df.to_csv(OUTPUT_CSV, index=False)
    with pd.ExcelWriter(OUTPUT_XLSX, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Bid List', index=False)
        wb = writer.book
        ws = writer.sheets['Bid List']
        format_bid_list(ws, df)
        build_summary(wb, df)
    validation = validate_workbook(OUTPUT_XLSX, df)
    total_opening = round(float(df['Opening Bid'].sum()), 2)
    cheapest = df.iloc[0]
    most_expensive = df.iloc[-1]
    print(f"output={OUTPUT_XLSX}")
    print(f"csv={OUTPUT_CSV}")
    print(f"count={len(df)}")
    print(f"cheapest={cheapest['Parcel ID']}|{cheapest['Address']}|{cheapest['Opening Bid']:.2f}")
    print(f"most_expensive={most_expensive['Parcel ID']}|{most_expensive['Address']}|{most_expensive['Opening Bid']:.2f}")
    print(f"total_opening={total_opening:.2f}")
    print(f"sorted={validation['sorted']}")
    print(f"duplicates_ok={validation['duplicates_ok']}")


if __name__ == '__main__':
    main()
