#!/usr/bin/env python3
from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from common_custom_33 import CLEANED_DIR, OUTPUT_DIR, ensure_dirs

FINAL_CSV = CLEANED_DIR / 'custom_33_final_ranked.csv'
FULL_XLSX = OUTPUT_DIR / 'custom_33_property_investment_ranking.xlsx'
TOP10_XLSX = OUTPUT_DIR / 'custom_top_10_priority_targets.xlsx'
DRIVEBY_XLSX = OUTPUT_DIR / 'custom_drive_by_list.xlsx'
WATCH_XLSX = OUTPUT_DIR / 'custom_high_risk_watchlist.xlsx'

TIER_COLOR = {
    'Tier 1 Priority': 'C6EFCE',
    'Tier 2 Strong Review': 'FFF2CC',
    'Tier 3 Watch': 'FCE4D6',
    'Tier 4 High Risk': 'F4CCCC',
}
RED_HIGHLIGHT = 'FFC7CE'
GRAY_HIGHLIGHT = 'D9D9D9'

REQUIRED_COLUMNS = [
    'final_rank', 'source_list', 'original_rank', 'parcel_id', 'property_address', 'city', 'state', 'bid_cost',
    'legal_description', 'property_type', 'confirmed_zip_code', 'neighborhood_or_area', 'geocode_status',
    'geocode_confidence', 'crime_risk_level', 'crime_risk_score', 'crime_confidence', 'median_household_income',
    'poverty_rate', 'unemployment_rate', 'median_home_value', 'vacancy_rate', 'economic_risk_level',
    'economic_strength_score', 'economic_confidence', 'rule_score', 'ai_score', 'final_score', 'final_tier',
    'recommendation', 'bid_strategy_note', 'reasoning_summary', 'crime_economic_summary', 'key_risks',
    'missing_information', 'next_due_diligence_step', 'confidence_level'
]


def write_df(df: pd.DataFrame, path):
    df[REQUIRED_COLUMNS].to_excel(path, index=False)
    format_book(path)


def format_book(path):
    wb = load_workbook(path)
    ws = wb.active
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='top')

    currency_cols = {'bid_cost', 'median_home_value', 'median_household_income'}
    pct_cols = {'poverty_rate', 'unemployment_rate', 'vacancy_rate'}

    for row_idx in range(2, ws.max_row + 1):
        tier = ws.cell(row_idx, headers['final_tier']).value
        row_fill = TIER_COLOR.get(tier)
        if row_fill:
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row_idx, col_idx).fill = PatternFill(fill_type='solid', start_color=row_fill, end_color=row_fill)
        if ws.cell(row_idx, headers['crime_risk_level']).value == 'High':
            for name in ['crime_risk_level', 'crime_risk_score', 'crime_confidence']:
                ws.cell(row_idx, headers[name]).fill = PatternFill(fill_type='solid', start_color=RED_HIGHLIGHT, end_color=RED_HIGHLIGHT)
        if ws.cell(row_idx, headers['economic_risk_level']).value == 'High':
            for name in ['economic_risk_level', 'economic_strength_score', 'economic_confidence']:
                ws.cell(row_idx, headers[name]).fill = PatternFill(fill_type='solid', start_color=RED_HIGHLIGHT, end_color=RED_HIGHLIGHT)
        if ws.cell(row_idx, headers['confidence_level']).value == 'Low':
            ws.cell(row_idx, headers['confidence_level']).fill = PatternFill(fill_type='solid', start_color=GRAY_HIGHLIGHT, end_color=GRAY_HIGHLIGHT)
        if ws.cell(row_idx, headers['geocode_status']).value == 'Failed':
            for name in ['geocode_status', 'geocode_confidence']:
                ws.cell(row_idx, headers[name]).fill = PatternFill(fill_type='solid', start_color=GRAY_HIGHLIGHT, end_color=GRAY_HIGHLIGHT)

    for name, idx in headers.items():
        letter = ws.cell(1, idx).column_letter
        max_len = max(len(str(ws.cell(r, idx).value or '')) for r in range(1, min(ws.max_row, 250) + 1))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, idx)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if name in currency_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '$#,##0.00'
            if name in pct_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '0.00%'

    wb.save(path)
    reopened = load_workbook(path)
    ws2 = reopened.active
    assert ws2.max_row > 1
    assert ws2.max_column > 1


def main() -> None:
    ensure_dirs()
    df = pd.read_csv(FINAL_CSV)
    write_df(df, FULL_XLSX)
    write_df(df.head(10).copy(), TOP10_XLSX)
    drive_by = df[df['recommendation'].isin(['Bid Candidate', 'Research First', 'Drive By'])].copy()
    if drive_by.empty:
        drive_by = df.head(10).copy()
    write_df(drive_by, DRIVEBY_XLSX)
    watch = df[(df['final_tier'].isin(['Tier 3 Watch', 'Tier 4 High Risk'])) | (df['recommendation'] == 'Avoid')].copy()
    if watch.empty:
        watch = df.tail(min(10, len(df))).copy()
    write_df(watch, WATCH_XLSX)
    print(f'excel_written={len(df)}')


if __name__ == '__main__':
    main()
