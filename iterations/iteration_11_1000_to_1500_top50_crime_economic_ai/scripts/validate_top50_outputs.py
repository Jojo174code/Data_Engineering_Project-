#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path
import subprocess

import pandas as pd
from openpyxl import load_workbook

ITERATION_ROOT = Path(__file__).resolve().parents[1]
CLEANED_DIR = ITERATION_ROOT / 'cleaned_data'
OUT_DIR = ITERATION_ROOT / 'output_excel'
LOG_DIR = ITERATION_ROOT / 'logs'

BASE = CLEANED_DIR / 'base_1000_to_1500_properties.csv'
GEO = CLEANED_DIR / 'geocoded_properties.csv'
AREA = CLEANED_DIR / 'property_area_intelligence.csv'
RANKED = CLEANED_DIR / 'crime_economic_ranked_properties.csv'
AI = CLEANED_DIR / 'ai_reviews_crime_economic.csv'
FINAL = CLEANED_DIR / 'final_top50_crime_economic_ai.csv'
TOP = OUT_DIR / 'top_50_properties_1000_to_1500_crime_economic_ai.xlsx'
FULL = OUT_DIR / 'full_ranked_properties_1000_to_1500_crime_economic_ai.xlsx'
RISK = OUT_DIR / 'high_risk_removed_properties_1000_to_1500.xlsx'
REPORT = LOG_DIR / 'validation_report.txt'
PREPARE_LOG = LOG_DIR / 'prepare_base_log.txt'
SOURCE_PREFERRED = '/home/ubuntu/OpenClaw_Context/iterations/iteration_06_final_robust_ai_due_diligence/cleaned_data/final_property_base.csv'
SOURCE_FALLBACK = '/home/ubuntu/OpenClaw_Context/iterations/iteration_04_development_intelligence_ai_review/cleaned_data/cleaned_auction_properties.csv'

ALLOWED_RISK = {'Low', 'Medium', 'High', 'Unknown'}
ALLOWED_CATEGORY = {'Top Candidate', 'Manual Review Candidate', 'Risky / Needs Verification', 'Avoid'}
ALLOWED_CONFIDENCE = {'High', 'Medium', 'Low'}


def workbook_ok(path: Path):
    wb = load_workbook(path)
    ws = wb.active
    assert ws.max_row > 1 and ws.max_column > 1


def reasoning_mentions_context(text: str) -> bool:
    lower = str(text).lower()
    has_crime = 'crime' in lower or 'crime risk' in lower
    has_econ = any(term in lower for term in ['income', 'poverty', 'unemployment', 'economic data is unknown', 'economic'])
    return has_crime and has_econ


def main():
    required = [BASE, GEO, AREA, RANKED, AI, FINAL, TOP, FULL, RISK]
    for path in required:
        if not path.exists():
            raise SystemExit(f'Missing required output: {path}')

    base = pd.read_csv(BASE)
    geo = pd.read_csv(GEO)
    area = pd.read_csv(AREA)
    ranked = pd.read_csv(RANKED)
    ai = pd.read_csv(AI)
    top = pd.read_csv(FINAL)
    full = pd.read_excel(FULL)

    assert len(base) > 0
    assert len(geo) > 0
    assert len(area) > 0
    assert len(ranked) > 0
    assert len(ai) > 0
    assert len(top) > 0
    assert len(top) <= 50

    for book in [TOP, FULL, RISK]:
        workbook_ok(book)

    assert top['bid_cost'].between(1000, 1500).all()
    assert top['parcel_id'].duplicated().sum() == 0
    assert set(top['crime_risk_level']).issubset(ALLOWED_RISK)
    assert set(top['economic_risk_level']).issubset(ALLOWED_RISK)
    assert set(top['final_category']).issubset(ALLOWED_CATEGORY)
    assert set(top['confidence_level']).issubset(ALLOWED_CONFIDENCE)
    assert top['reasoning_summary'].apply(reasoning_mentions_context).all()

    status = subprocess.check_output(['git', 'status', '--short'], cwd='/home/ubuntu/OpenClaw_Context', text=True)
    assert '.env' not in status
    assert '.venv' not in status

    if PREPARE_LOG.exists():
        first_line = PREPARE_LOG.read_text().splitlines()[0]
        source_used = first_line.split(': ', 1)[1] if ': ' in first_line else SOURCE_FALLBACK
    else:
        source_used = SOURCE_PREFERRED if Path(SOURCE_PREFERRED).exists() else SOURCE_FALLBACK
    geocoded_success = int((geo['geocode_status'] != 'Failed').sum())
    crime_coverage = int((area['crime_source'] != 'Unknown.').sum())
    economic_coverage = int((area['economic_source'] != 'Unknown.').sum())
    top_candidate_count = int((top['final_category'] == 'Top Candidate').sum())
    manual_count = int((top['final_category'] == 'Manual Review Candidate').sum())
    risky_count = int((top['final_category'] == 'Risky / Needs Verification').sum())
    avoid_count = int((top['final_category'] == 'Avoid').sum())
    high_crime_count = int((top['crime_risk_level'] == 'High').sum())
    unknown_crime_count = int((top['crime_risk_level'] == 'Unknown').sum())
    high_economic_count = int((top['economic_risk_level'] == 'High').sum())
    unknown_economic_count = int((top['economic_risk_level'] == 'Unknown').sum())

    report_lines = [
        f'source file used: {source_used}',
        f'total properties in $1,000-$1,500 range: {len(base)}',
        f'total geocoded successfully: {geocoded_success}',
        f'total with crime data: {crime_coverage}',
        f'total with economic data: {economic_coverage}',
        f'top 50 count: {len(top)}',
        f'Top Candidate count: {top_candidate_count}',
        f'Manual Review Candidate count: {manual_count}',
        f'Risky / Needs Verification count: {risky_count}',
        f'Avoid count: {avoid_count}',
        f'High crime risk count: {high_crime_count}',
        f'Unknown crime risk count: {unknown_crime_count}',
        f'High economic risk count: {high_economic_count}',
        f'Unknown economic risk count: {unknown_economic_count}',
        'validation PASS',
    ]
    REPORT.write_text('\n'.join(report_lines))
    print('PASS')


if __name__ == '__main__':
    main()
