#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ITERATION_ROOT = Path(__file__).resolve().parents[1]
CLEANED_DIR = ITERATION_ROOT / 'cleaned_data'
OUT_DIR = ITERATION_ROOT / 'output_excel'
RANKED_CSV = CLEANED_DIR / 'crime_economic_ranked_properties.csv'
AI_CSV = CLEANED_DIR / 'ai_reviews_crime_economic.csv'
FINAL_TOP50_CSV = CLEANED_DIR / 'final_top50_crime_economic_ai.csv'
TOP_XLSX = OUT_DIR / 'top_50_properties_1000_to_1500_crime_economic_ai.xlsx'
FULL_XLSX = OUT_DIR / 'full_ranked_properties_1000_to_1500_crime_economic_ai.xlsx'
RISK_XLSX = OUT_DIR / 'high_risk_removed_properties_1000_to_1500.xlsx'

CATEGORY_COLOR = {
    'Top Candidate': 'C6EFCE',
    'Manual Review Candidate': 'FFF2CC',
    'Risky / Needs Verification': 'FCE4D6',
    'Avoid': 'F4CCCC',
}
RED_HIGHLIGHT = 'FFC7CE'
GRAY_HIGHLIGHT = 'D9D9D9'

TOP_COLUMNS = [
    'final_rank', 'parcel_id', 'owner_name', 'property_address', 'city', 'state', 'zip_code', 'confirmed_zip_code',
    'bid_cost', 'legal_description', 'property_type', 'geocode_status', 'geocode_confidence', 'neighborhood_or_area',
    'crime_risk_level', 'crime_risk_score', 'crime_confidence', 'crime_source', 'crime_data_date_range',
    'median_household_income', 'poverty_rate', 'unemployment_rate', 'median_home_value', 'median_gross_rent',
    'vacancy_rate', 'owner_occupied_rate', 'economic_risk_level', 'economic_strength_score', 'economic_confidence',
    'economic_source', 'property_clarity_score', 'residential_likelihood_score', 'pre_ai_final_score',
    'ai_area_adjusted_score', 'final_score', 'final_category', 'recommendation', 'reasoning_summary',
    'crime_economic_summary', 'key_risks', 'missing_information', 'next_due_diligence_step', 'confidence_level',
    'raw_text'
]


def clean_text(value, default='Unknown.'):
    if pd.isna(value):
        return default
    text = str(value).strip()
    if not text or text.lower() == 'nan':
        return default
    return text


def num(value, fallback=None):
    try:
        return float(value)
    except Exception:
        return fallback


def category_from_score(row: pd.Series) -> str:
    score = float(row['final_score'])
    if clean_text(row['crime_risk_level']) == 'High' and score < 7.9:
        return 'Risky / Needs Verification' if score >= 4.0 else 'Avoid'
    if clean_text(row['economic_risk_level']) == 'High' and score < 7.9:
        return 'Risky / Needs Verification' if score >= 4.0 else 'Avoid'
    if clean_text(row['property_address']) in {'Unknown.', 'Unknown', 'ADDRESS UNKNOWN'} and clean_text(row['confidence_level']) == 'Low':
        return 'Risky / Needs Verification' if score >= 4.0 else 'Avoid'
    if score >= 7.5 and clean_text(row['crime_risk_level']) != 'High' and clean_text(row['economic_risk_level']) != 'High' and clean_text(row['confidence_level']) in {'Medium', 'High'}:
        return 'Top Candidate'
    if score >= 5.5:
        return 'Manual Review Candidate'
    if score >= 4.0:
        return 'Risky / Needs Verification'
    return 'Avoid'


def sort_priority(df: pd.DataFrame) -> pd.DataFrame:
    category_order = {'Top Candidate': 0, 'Manual Review Candidate': 1, 'Risky / Needs Verification': 2, 'Avoid': 3}
    risk_order = {'Low': 0, 'Medium': 1, 'Unknown': 2, 'Unknown.': 2, 'High': 3}
    return df.assign(
        _cat=df['final_category'].map(category_order).fillna(9),
        _crime=df['crime_risk_level'].map(risk_order).fillna(9),
        _econ=df['economic_risk_level'].map(risk_order).fillna(9),
    ).sort_values(['_cat', 'final_score', '_crime', '_econ', 'bid_cost'], ascending=[True, False, True, True, True]).drop(columns=['_cat', '_crime', '_econ']).reset_index(drop=True)


def format_book(path: Path, category_col='final_category'):
    wb = load_workbook(path)
    ws = wb.active
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='top')

    currency_cols = {'bid_cost', 'median_home_value', 'median_gross_rent', 'median_household_income'}
    pct_cols = {'poverty_rate', 'unemployment_rate', 'vacancy_rate', 'owner_occupied_rate'}

    for row_idx in range(2, ws.max_row + 1):
        category = ws.cell(row_idx, headers.get(category_col, 1)).value if category_col in headers else None
        row_fill = CATEGORY_COLOR.get(category)
        if row_fill:
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row_idx, col_idx).fill = PatternFill(fill_type='solid', start_color=row_fill, end_color=row_fill)

        if 'crime_risk_level' in headers and ws.cell(row_idx, headers['crime_risk_level']).value == 'High':
            for name in ['crime_risk_level', 'crime_risk_score', 'crime_confidence']:
                if name in headers:
                    ws.cell(row_idx, headers[name]).fill = PatternFill(fill_type='solid', start_color=RED_HIGHLIGHT, end_color=RED_HIGHLIGHT)
        if 'economic_risk_level' in headers and ws.cell(row_idx, headers['economic_risk_level']).value == 'High':
            for name in ['economic_risk_level', 'economic_strength_score', 'economic_confidence']:
                if name in headers:
                    ws.cell(row_idx, headers[name]).fill = PatternFill(fill_type='solid', start_color=RED_HIGHLIGHT, end_color=RED_HIGHLIGHT)
        for name in ['crime_confidence', 'economic_confidence']:
            if name in headers and ws.cell(row_idx, headers[name]).value in {'Unknown', 'Unknown.', 'Low'}:
                ws.cell(row_idx, headers[name]).fill = PatternFill(fill_type='solid', start_color=GRAY_HIGHLIGHT, end_color=GRAY_HIGHLIGHT)

    for name, idx in headers.items():
        letter = ws.cell(1, idx).column_letter
        max_len = max(len(str(ws.cell(r, idx).value or '')) for r in range(1, min(ws.max_row, 250) + 1))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, idx)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if name in currency_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '$#,##0.00'
            if name in pct_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '0.00%'

    wb.save(path)
    reopened = load_workbook(path)
    ws2 = reopened.active
    assert ws2.max_row > 1 and ws2.max_column > 1


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ranked = pd.read_csv(RANKED_CSV)
    ai = pd.read_csv(AI_CSV)
    merged = ranked.merge(ai, on='parcel_id', how='left')

    merged['ai_area_adjusted_score'] = merged['ai_area_adjusted_score'].apply(lambda v: round(float(v), 2))
    merged['final_score'] = ((merged['pre_ai_final_score'] * 0.60) + (merged['ai_area_adjusted_score'] * 0.40)).round(2)

    for idx, row in merged.iterrows():
        if clean_text(row['crime_risk_level']) == 'High':
            merged.at[idx, 'final_score'] -= 1.0
        if clean_text(row['economic_risk_level']) == 'High':
            merged.at[idx, 'final_score'] -= 0.75
        if clean_text(row['geocode_confidence']) == 'Low':
            merged.at[idx, 'final_score'] -= 0.75
        if clean_text(row['crime_confidence']) == 'Low':
            merged.at[idx, 'final_score'] -= 0.5
        if clean_text(row['economic_confidence']) == 'Low':
            merged.at[idx, 'final_score'] -= 0.5
        if clean_text(row['property_address']) in {'Unknown.', 'Unknown', 'ADDRESS UNKNOWN'}:
            merged.at[idx, 'final_score'] -= 0.5
        if clean_text(row['parcel_id']) in {'Unknown.', 'Unknown'}:
            merged.at[idx, 'final_score'] -= 0.5
        if clean_text(row['property_type']).upper().startswith('C'):
            merged.at[idx, 'final_score'] -= 0.5
        if clean_text(row['crime_risk_level']) == 'Low' and clean_text(row['economic_risk_level']) == 'Low':
            merged.at[idx, 'final_score'] += 0.25
        if clean_text(row['geocode_confidence']) == 'High' and clean_text(row['data_confidence']) == 'High':
            merged.at[idx, 'final_score'] += 0.25

    merged['final_score'] = merged['final_score'].apply(lambda v: round(max(1.0, min(10.0, float(v))), 2))
    merged['final_category'] = merged.apply(category_from_score, axis=1)
    merged = sort_priority(merged)

    top50 = merged.head(50).copy()
    top50['final_rank'] = range(1, len(top50) + 1)
    final_top50 = top50[TOP_COLUMNS].copy()
    final_top50.to_csv(FINAL_TOP50_CSV, index=False)

    risk_only = merged[~merged['parcel_id'].isin(top50['parcel_id']) & merged['final_category'].isin(['Risky / Needs Verification', 'Avoid'])].copy()
    if risk_only.empty:
        risk_only = merged[merged['final_category'].isin(['Risky / Needs Verification', 'Avoid'])].copy()
    if risk_only.empty:
        risk_only = merged.tail(min(3, len(merged))).copy()

    full_columns = ['parcel_id', 'owner_name', 'property_address', 'city', 'state', 'zip_code', 'confirmed_zip_code', 'bid_cost', 'legal_description', 'property_type', 'geocode_status', 'geocode_confidence', 'neighborhood_or_area', 'crime_risk_level', 'crime_risk_score', 'crime_confidence', 'crime_source', 'crime_data_date_range', 'median_household_income', 'poverty_rate', 'unemployment_rate', 'median_home_value', 'median_gross_rent', 'vacancy_rate', 'owner_occupied_rate', 'economic_risk_level', 'economic_strength_score', 'economic_confidence', 'economic_source', 'property_clarity_score', 'geocoding_score', 'residential_likelihood_score', 'crime_score', 'bid_price_score', 'data_confidence_score', 'pre_ai_final_score', 'pre_ai_category', 'ai_area_adjusted_score', 'ai_area_adjusted_category', 'final_score', 'final_category', 'recommendation', 'reasoning_summary', 'crime_economic_summary', 'key_risks', 'missing_information', 'next_due_diligence_step', 'confidence_level', 'raw_text']

    final_top50.to_excel(TOP_XLSX, index=False)
    merged[full_columns].to_excel(FULL_XLSX, index=False)
    risk_only[full_columns].to_excel(RISK_XLSX, index=False)

    format_book(TOP_XLSX)
    format_book(FULL_XLSX)
    format_book(RISK_XLSX)
    print(f'top50_written={len(final_top50)}')


if __name__ == '__main__':
    main()
