from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ITERATION_ROOT = Path(__file__).resolve().parents[1]
INPUT_PDF = ITERATION_ROOT / 'input' / '2026_tulsa_auction_list.pdf'
CLEANED_DIR = ITERATION_ROOT / 'cleaned_data'
OUTPUT_DIR = ITERATION_ROOT / 'output_excel'
LOG_DIR = ITERATION_ROOT / 'logs'

ILLEGAL_CHAR_RE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
CITY_SUFFIX_RE = re.compile(r'\s+((?:CITY|TOWN) OF [A-Z ]+|UNINCORPORATED|BROKEN ARROW|SAND SPRINGS|OWASSO|BIXBY|JENKS|SAPULPA|COLLINSVILLE|SKIATOOK)$', re.IGNORECASE)
ADDRESS_RE = re.compile(r'^(ADDRESS UNKNOWN|\d+\s+.*)$', re.IGNORECASE)
ALLOWED_CATEGORIES = {'Good Investment', 'Strong Manual Review Candidate', 'Mid Investment', 'Bad Investment'}
ALLOWED_RECOMMENDATIONS = {'Bid Candidate', 'Research First', 'Drive By', 'Watch', 'Avoid'}
ALLOWED_CONFIDENCE = {'High', 'Medium', 'Low'}
ALLOWED_MATCH_STRENGTH = {'Strong Match', 'Possible Match', 'No Clear Match', 'Unknown'}
ALLOWED_MATCH_QUALITY = {'High', 'Medium', 'Low', 'Invalid Weak Match'}
STREET_TYPE_WORDS = {'ST', 'AVE', 'AV', 'PL', 'RD', 'DR', 'BLVD', 'BV', 'HWY', 'PKWY', 'LN', 'CT', 'TER', 'WAY', 'N', 'S', 'E', 'W'}
WEAK_LOCATION_TOKENS = {
    'TULSA', 'OKLAHOMA', 'COUNTY', 'CITY', 'NORTH', 'SOUTH', 'EAST', 'WEST',
    'INDUSTRIAL', 'PARK', 'CORRIDOR', 'CITYWIDE', 'METRO', 'REGIONAL', 'MULTIPLE', 'URBAN', 'CORE',
    'UNPLATTED', 'ADDN', 'ADDITION', 'SUB', 'SUBD', 'RESUB', 'BLK', 'LOT', 'LTS', 'LT'
}


def clean_text(value) -> str:
    if value is None:
        return ''
    return ILLEGAL_CHAR_RE.sub('', str(value)).strip()


def normalize_city(value: str) -> str:
    text = clean_text(value)
    if not text:
        return 'Unknown'
    replacements = {
        'CITY OF TULSA': 'Tulsa',
        'TOWN OF JENKS': 'Jenks',
        'TOWN OF BIXBY': 'Bixby',
        'CITY OF OWASSO': 'Owasso',
        'CITY OF SAND SPRINGS': 'Sand Springs',
        'CITY OF BROKEN ARROW': 'Broken Arrow',
        'CITY OF COLLINSVILLE': 'Collinsville',
        'CITY OF SKIATOOK': 'Skiatook',
    }
    return replacements.get(text.upper(), text.title())


def parse_money(value):
    text = clean_text(value).replace('$', '').replace(',', '')
    if not text or text == 'Unknown':
        return None
    try:
        return float(text)
    except ValueError:
        return None


def split_legal_city(line: str):
    line = clean_text(line)
    m = CITY_SUFFIX_RE.search(line)
    if not m:
        return line, 'Unknown'
    city_raw = clean_text(m.group(1))
    legal = clean_text(line[:m.start()])
    return legal or 'Unknown', normalize_city(city_raw)


def looks_residential_address(address: str) -> bool:
    return bool(ADDRESS_RE.match(clean_text(address))) and clean_text(address).upper() != 'ADDRESS UNKNOWN'


def infer_extraction_confidence(parcel_id: str, address: str, legal_description: str, bid_cost) -> tuple[str, str]:
    score = 0
    notes = []
    if clean_text(parcel_id) and clean_text(parcel_id) != 'Unknown':
        score += 1
    else:
        notes.append('missing parcel_id')
    if looks_residential_address(address):
        score += 1
    else:
        notes.append('partial or unknown address')
    if clean_text(legal_description) and clean_text(legal_description) != 'Unknown':
        score += 1
    else:
        notes.append('missing legal description')
    if bid_cost is not None:
        score += 1
    else:
        notes.append('missing bid_cost')
    if score >= 4:
        return 'High', '; '.join(notes) or 'parsed cleanly'
    if score >= 2:
        return 'Medium', '; '.join(notes) or 'usable row'
    return 'Low', '; '.join(notes) or 'low confidence row'


def tokenize_street_text(value: str, *, remove_weak: bool = True) -> set[str]:
    text = clean_text(value).upper().replace(',', ' ')
    parts = re.split(r'[^A-Z0-9]+', text)
    tokens = set()
    for part in parts:
        if not part:
            continue
        if part in STREET_TYPE_WORDS:
            continue
        if part.isdigit() and part != '66':
            continue
        if len(part) < 3 and part != '66':
            continue
        if remove_weak and part in WEAK_LOCATION_TOKENS:
            continue
        tokens.add(part)
    return tokens


def parse_zip_list(value: str) -> set[str]:
    text = clean_text(value)
    return set(re.findall(r'\b\d{5}\b', text))


def write_log(path: Path, lines: list[str]):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text('\n'.join(lines) + '\n', encoding='utf-8')


def write_dataframe_to_excel(df, path: Path, sheet_name: str, currency_columns: Iterable[str] | None = None,
                             wrap_columns: Iterable[str] | None = None,
                             category_fill_column: str | None = None,
                             category_fills: dict[str, str] | None = None,
                             highlight_column: str | None = None,
                             highlight_values: set[str] | None = None,
                             highlight_fill: str = 'FFF2CC'):
    import pandas as pd

    if df.empty:
        raise ValueError(f'Refusing to write empty workbook: {path}')
    path.parent.mkdir(parents=True, exist_ok=True)

    currency_columns = set(currency_columns or [])
    wrap_columns = set(wrap_columns or [])

    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.book[sheet_name]
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        header_font = Font(bold=True)
        wrap = Alignment(wrap_text=True, vertical='top')
        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = wrap
        header_map = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                header = ws.cell(row=1, column=cell.column).value
                if header in currency_columns and isinstance(cell.value, (int, float)):
                    cell.number_format = '$#,##0.00'
                if header in wrap_columns:
                    cell.alignment = wrap
        for r in range(2, ws.max_row + 1):
            row_fill = None
            if category_fill_column and category_fills and category_fill_column in header_map:
                category = ws.cell(row=r, column=header_map[category_fill_column]).value
                color = category_fills.get(str(category))
                if color:
                    row_fill = PatternFill(fill_type='solid', fgColor=color)
            if row_fill:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).fill = row_fill
            if highlight_column and highlight_values and highlight_column in header_map:
                value = str(ws.cell(row=r, column=header_map[highlight_column]).value)
                if value in highlight_values:
                    fill = PatternFill(fill_type='solid', fgColor=highlight_fill)
                    for c in range(1, ws.max_column + 1):
                        if ws.cell(row=r, column=c).fill.fill_type is None:
                            ws.cell(row=r, column=c).fill = fill
        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for cell in ws[col_letter]:
                value = '' if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 55)

    return validate_workbook(path)


def validate_workbook(path: Path):
    wb = load_workbook(path)
    if not wb.sheetnames:
        raise ValueError(f'Workbook has no sheets: {path}')
    ws = wb[wb.sheetnames[0]]
    if ws.max_row <= 1 or ws.max_column <= 1:
        raise ValueError(f'Workbook appears empty: {path}')
    return {
        'sheet_names': wb.sheetnames,
        'max_row': ws.max_row,
        'max_column': ws.max_column,
        'top_rows': [[cell.value for cell in row] for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row))],
    }
