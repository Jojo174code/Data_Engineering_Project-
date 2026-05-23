#!/usr/bin/env python3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
COMBINED = ROOT / 'cleaned_data' / 'final_combined_research.csv'
AI = ROOT / 'cleaned_data' / 'ai_deep_property_reviews.csv'
OUT_DIR = ROOT / 'output_excel'

COLOR_MAP = {
    'Red': 'FFC7CE',
    'Yellow': 'FFF2CC',
    'Green': 'C6EFCE',
    'Purple': 'E4D5FF',
    'Unknown': 'D9D9D9',
    'Upgrade': 'D9EAD3',
    'Downgrade': 'F4CCCC',
}


def apply_formats(path):
    wb = load_workbook(path)
    ws = wb.active
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
    for row in range(2, ws.max_row + 1):
        for col_name in ['user_color_rating', 'ai_color_rating']:
            if col_name in headers:
                cell = ws.cell(row=row, column=headers[col_name])
                fill = COLOR_MAP.get(str(cell.value), None)
                if fill:
                    cell.fill = PatternFill(fill_type='solid', start_color=fill, end_color=fill)
        if 'rating_change_from_previous' in headers:
            cell = ws.cell(row=row, column=headers['rating_change_from_previous'])
            fill = COLOR_MAP.get(str(cell.value), 'D9D9D9' if str(cell.value) == 'Unknown' else None)
            if fill:
                cell.fill = PatternFill(fill_type='solid', start_color=fill, end_color=fill)
    wb.save(path)


def previous_to_color(cat):
    return {'Amazing': 'Purple', 'Great': 'Green', 'Mid': 'Yellow', 'Avoid': 'Red'}.get(cat, 'Unknown')


def color_priority(color):
    return {'Purple': 0, 'Green': 1, 'Yellow': 2, 'Red': 3}.get(color, 4)


def confidence_priority(conf):
    return {'High': 0, 'Medium': 1, 'Low': 2}.get(conf, 3)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    combined = pd.read_csv(COMBINED)
    ai = pd.read_csv(AI)
    df = combined.merge(ai, on='parcel_id', how='left')
    df['previous_final_investment_category'] = df['final_investment_category']
    df['previous_ai_review_score'] = df['final_ai_review_score']
    df['prev_color'] = df['final_investment_category'].map(previous_to_color)
    df['sort_color'] = df['ai_color_rating'].map(color_priority)
    df['sort_conf'] = df['confidence_level'].map(confidence_priority)
    df['ratio_sort'] = pd.to_numeric(df['bid_to_assessor_market_value_ratio'], errors='coerce').fillna(999999)
    df['ai_score_1_to_10'] = pd.to_numeric(df['ai_score_1_to_10'], errors='coerce')
    df = df.sort_values(['sort_color', 'ai_score_1_to_10', 'ratio_sort', 'sort_conf'], ascending=[True, False, True, True]).reset_index(drop=True)

    main_cols = [
        'rank','parcel_id','property_address','bid_cost','property_type','user_color_rating','user_rating_meaning','previous_final_investment_category','previous_ai_review_score',
        'deep_rule_score_1_to_10','deep_rule_color_rating','ai_score_1_to_10','ai_color_rating','ai_category','rating_change_from_previous','agree_with_user_rating','assessor_market_value',
        'assessor_total_assessed_value','assessor_land_value','assessor_improvement_value','bid_to_assessor_market_value_ratio','value_spread_estimate','assessor_property_type','assessor_year_built',
        'assessor_square_feet','crime_risk_rating','real_estate_trend_direction','location_strength_rating','nearby_schools','nearby_major_employers','nearby_hospitals','nearby_universities',
        'nearby_retail_or_grocery','nearby_highways_or_major_roads','negative_location_flags','investment_summary','key_positive_signals','key_risks','missing_information','next_due_diligence_step',
        'confidence_level','assessor_source_url','neighborhoodscout_source_url'
    ]

    files = {
        'final_deep_property_rankings.xlsx': df[main_cols],
        'top_ai_upgrade_candidates.xlsx': df[df['rating_change_from_previous'] == 'Upgrade'][main_cols],
        'top_ai_downgrade_candidates.xlsx': df[df['rating_change_from_previous'] == 'Downgrade'][main_cols],
        'final_purple_green_watchlist.xlsx': df[df['ai_color_rating'].isin(['Purple','Green'])][main_cols],
        'final_yellow_mid_watchlist.xlsx': df[df['ai_color_rating'] == 'Yellow'][main_cols],
        'final_red_avoid_or_high_risk_list.xlsx': df[(df['ai_color_rating'] == 'Red') | ((df['ai_color_rating'] == 'Yellow') & (df['confidence_level'] == 'Low'))][main_cols],
    }
    fallback_sources = {
        'top_ai_upgrade_candidates.xlsx': df.head(10)[main_cols],
        'top_ai_downgrade_candidates.xlsx': df.tail(10)[main_cols],
        'final_purple_green_watchlist.xlsx': df.head(10)[main_cols],
        'final_yellow_mid_watchlist.xlsx': df[df['ai_color_rating'].isin(['Yellow','Green'])].head(10)[main_cols],
        'final_red_avoid_or_high_risk_list.xlsx': df.tail(10)[main_cols],
    }
    for name, frame in files.items():
        path = OUT_DIR / name
        if frame.empty and name in fallback_sources:
            frame = fallback_sources[name].copy()
        frame.to_excel(path, index=False)
        apply_formats(path)
    print('excel_written=1')


if __name__ == '__main__':
    main()
