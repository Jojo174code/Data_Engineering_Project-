#!/usr/bin/env python3
from pathlib import Path
import subprocess

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FILTERED = ROOT / 'cleaned_data' / 'filtered_properties_1500_to_2000.csv'
RANKED = ROOT / 'cleaned_data' / 'ranked_properties_1500_to_2000.csv'
AI = ROOT / 'cleaned_data' / 'ai_reviews_1500_to_2000.csv'
FX = ROOT / 'output_excel' / 'filtered_properties_1500_to_2000.xlsx'
TOP = ROOT / 'output_excel' / 'top_50_properties_1500_to_2000.xlsx'
RISK = ROOT / 'output_excel' / 'high_risk_properties_1500_to_2000.xlsx'
LOG = ROOT / 'logs' / 'validation_report.txt'
SOURCE = '/home/ubuntu/OpenClaw_Context/iterations/iteration_04_development_intelligence_ai_review/cleaned_data/cleaned_auction_properties.csv'

ALLOWED = {'Top Candidate', 'Manual Review Candidate', 'Risky / Needs Verification', 'Avoid'}


def main():
    for p in [FILTERED, RANKED, AI, FX, TOP, RISK]:
        if not p.exists():
            raise SystemExit(f'Missing required output: {p}')

    filtered = pd.read_csv(FILTERED)
    ranked = pd.read_csv(RANKED)
    ai = pd.read_csv(AI)
    top = pd.read_excel(TOP)

    assert len(filtered) > 0
    assert len(ranked) > 0
    assert len(ai) > 0

    for p in [FX, TOP, RISK]:
        wb = load_workbook(p)
        ws = wb.active
        assert ws.max_row > 1 and ws.max_column > 1

    assert len(top) <= 50
    assert top['bid_cost'].between(1500, 2000).all()
    assert top['parcel_id'].duplicated().sum() == 0
    assert set(top['final_category']).issubset(ALLOWED)

    status = subprocess.check_output(['git', 'status', '--short'], text=True)
    assert '.env' not in status
    assert '.venv' not in status

    report = [
        f'source file used: {SOURCE}',
        f'total filtered properties: {len(filtered)}',
        f'total ranked properties: {len(ranked)}',
        f'total AI-reviewed properties: {len(ai)}',
        f'top 50 count: {len(top)}',
        f'high-risk count: {len(pd.read_excel(RISK))}',
        f'average bid cost in top 50: {top["bid_cost"].mean():.2f}',
        f'lowest bid in top 50: {top["bid_cost"].min():.2f}',
        f'highest bid in top 50: {top["bid_cost"].max():.2f}',
        f'Top Candidate count: {(top["final_category"] == "Top Candidate").sum()}',
        f'Manual Review Candidate count: {(top["final_category"] == "Manual Review Candidate").sum()}',
        f'Risky / Needs Verification count: {(top["final_category"] == "Risky / Needs Verification").sum()}',
        f'Avoid count: {(top["final_category"] == "Avoid").sum()}',
        'validation PASS'
    ]
    LOG.write_text('\n'.join(report))
    print('PASS')


if __name__ == '__main__':
    main()
