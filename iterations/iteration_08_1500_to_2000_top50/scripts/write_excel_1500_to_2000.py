#!/usr/bin/env python3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
FILTERED = ROOT / 'cleaned_data' / 'filtered_properties_1500_to_2000.csv'
RANKED = ROOT / 'cleaned_data' / 'ranked_properties_1500_to_2000.csv'
AI = ROOT / 'cleaned_data' / 'ai_reviews_1500_to_2000.csv'
OUT = ROOT / 'output_excel'

COLOR = {
    'Top Candidate': 'C6EFCE',
    'Manual Review Candidate': 'FFF2CC',
    'Risky / Needs Verification': 'FCE4D6',
    'Avoid': 'F4CCCC',
}


def format_book(path, category_col=None):
    wb = load_workbook(path)
    ws = wb.active
    ws.freeze_panes = 'A2'
    headers = {c.value: i+1 for i, c in enumerate(ws[1])}
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    if 'bid_cost' in headers:
        for r in range(2, ws.max_row+1):
            ws.cell(r, headers['bid_cost']).number_format = '$#,##0.00'
    if category_col and category_col in headers:
        for r in range(2, ws.max_row+1):
            cat = ws.cell(r, headers[category_col]).value
            fill = COLOR.get(cat)
            if fill:
                for c in range(1, ws.max_column+1):
                    ws.cell(r, c).fill = PatternFill(fill_type='solid', start_color=fill, end_color=fill)
    for col in ws.columns:
        length = max(len(str(cell.value or '')) for cell in col[:100])
        ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 12), 40)
    wb.save(path)
    wb2 = load_workbook(path)
    ws2 = wb2.active
    assert ws2.max_row > 1 and ws2.max_column > 1


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    filtered = pd.read_csv(FILTERED)
    ranked = pd.read_csv(RANKED)
    ai = pd.read_csv(AI)
    merged = ranked.merge(ai, on='parcel_id', how='left')
    merged['final_score'] = merged['ai_score'].fillna(merged['rank_score'])
    merged['final_category'] = merged['ai_category'].fillna(merged['rank_category'])
    merged = merged.sort_values(['final_score', 'bid_cost'], ascending=[False, True]).reset_index(drop=True)
    top50 = merged.head(50).copy()
    top50['final_rank'] = range(1, len(top50)+1)
    high_risk = merged[~merged['parcel_id'].isin(top50['parcel_id']) | merged['final_category'].isin(['Risky / Needs Verification', 'Avoid'])].copy()
    if high_risk.empty:
        high_risk = merged.tail(min(10, len(merged))).copy()

    filtered_path = OUT / 'filtered_properties_1500_to_2000.xlsx'
    top50_path = OUT / 'top_50_properties_1500_to_2000.xlsx'
    risk_path = OUT / 'high_risk_properties_1500_to_2000.xlsx'

    filtered.to_excel(filtered_path, index=False)

    top_cols = [
        'final_rank','parcel_id','owner_name','property_address','city','state','zip_code','bid_cost','legal_description','property_type','source_page','residential_likelihood',
        'rank_score','ai_score','final_score','final_category','recommendation','reasoning_summary','key_risks','missing_information','next_due_diligence_step','confidence_level','raw_text'
    ]
    top50[top_cols].to_excel(top50_path, index=False)
    high_risk[top_cols[1:]].to_excel(risk_path, index=False)

    format_book(filtered_path)
    format_book(top50_path, 'final_category')
    format_book(risk_path, 'final_category')
    print('excel_outputs_written=1')


if __name__ == '__main__':
    main()
